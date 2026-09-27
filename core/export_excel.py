"""Export des Schichtplans im Aufbau des Regieplans (ein Tabellenblatt pro Monat).

Spalten wie im bestehenden Regieplan des Theaters:
  Datum | Tag | Uhrzeit | Vorstellungen/Bühne | Technik | ASL | Einlass | Tageskasse |
  Abendkasse | Proben (+Sperrtermine) | Sonstiges

Jeder Tag ist ein Block aus mindestens 3 Zeilen (mehr, wenn mehr Namen nötig sind),
getrennt durch eine dickere Linie. Hellblau = von der App ausgefüllt (inkl. Technik), weiß = frei für
ASL, Proben und Sonstiges. Unbesetzte Plätze erscheinen rot als „offen".
"""
import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import db
from .calendar_utils import WEEKDAYS, month_label, month_weeks, parse_month

HEADERS = ["Datum", "Tag", "Uhrzeit", "Vorstellungen/Bühne", "Technik", "ASL", "Einlass",
           "Tageskasse", "Abendkasse", "Proben (+Sperrtermine)", "Sonstiges"]
COL = {h: i + 1 for i, h in enumerate(HEADERS)}
APP_COLUMNS = {"Datum", "Tag", "Uhrzeit", "Vorstellungen/Bühne", "Technik", "Einlass", "Tageskasse", "Abendkasse"}
TYPE_COLUMN = {"EINLASS": "Einlass", "ABENDKASSE": "Abendkasse", "TECHNIK": "Technik"}
EVENT_COLUMNS = ("Technik", "Einlass", "Abendkasse")
# Spaltenbreiten: gleicher Spaltentyp = gleiche Breite, mindestens 13
WIDTHS = {"Datum": 13, "Tag": 13, "Uhrzeit": 13, "Vorstellungen/Bühne": 34, "Technik": 18, "ASL": 18,
          "Einlass": 18, "Tageskasse": 18, "Abendkasse": 18, "Proben (+Sperrtermine)": 34, "Sonstiges": 34}
HEADER_HEIGHT, ROW_HEIGHT, MIN_ROWS_PER_DAY = 42, 15, 3

FONT = Font(name="Arial", size=11, color="000000")
FONT_BOLD = Font(name="Arial", size=11, bold=True, color="000000")
FONT_OPEN = Font(name="Arial", size=11, bold=True, color="C00000")
FONT_LEGEND = Font(name="Arial", size=9, italic=True, color="000000")
FILL_HEADER = PatternFill("solid", fgColor="C9DAF8")
FILL_APP = PatternFill("solid", fgColor="E8F0FE")
FILL_FREE = PatternFill("solid", fgColor="FFFFFF")
THIN, MEDIUM = Side(style="thin", color="000000"), Side(style="medium", color="000000")
ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_HEADER = Alignment(horizontal="left", vertical="top", wrap_text=True)
OPEN = "offen"


def assignment_names(month: str, trial: list[tuple[int, int]] | None = None) -> dict[int, list[str]]:
    """Namen je Schicht – aus dem gespeicherten Plan oder aus einer Probe-Zuteilung (shift_id, user_id)."""
    names: dict[int, list[str]] = defaultdict(list)
    if trial is None:
        for a in db.get_assignments(month):
            names[a["shift_id"]].append(a["name"])
    else:
        user_name = {u["id"]: u["name"] for u in db.list_users()}
        for sid, uid in trial:
            names[sid].append(user_name.get(uid, "?"))
    return names


def _day_data(month: str, trial: list[tuple[int, int]] | None = None) -> dict[str, dict]:
    """Sammelt pro Tag: Vorstellungen (mit Einlass/Abendkasse) und Tageskasse."""
    names = assignment_names(month, trial)
    perfs = {p["id"]: p for p in db.list_performances(month)}
    days: dict[str, dict] = defaultdict(lambda: {"events": {}, "tageskasse": []})

    for p in perfs.values():                                   # auch Vorstellungen ohne Schichten zeigen
        days[p["date"]]["events"][("p", p["id"])] = {"time": p["time"], "title": p["title"],
                                                     **{c: [] for c in EVENT_COLUMNS}}
    for s in db.list_shifts(month):
        staffed = sorted(names.get(s["id"], []))
        staffed += [OPEN] * max(0, s["required"] - len(staffed))
        if s["shift_type"] == "TAGESKASSE":
            days[s["date"]]["tageskasse"] += staffed
            continue
        pid = s.get("performance_id")
        if pid in perfs:
            p = perfs[pid]
            ev = days[p["date"]]["events"][("p", pid)]
        else:                                                  # Schicht ohne verknüpfte Vorstellung
            key = ("t", s["title"] or "")
            ev = days[s["date"]]["events"].setdefault(key, {"time": "", "title": s["title"] or "",
                                                            **{c: [] for c in EVENT_COLUMNS}})
        ev[TYPE_COLUMN[s["shift_type"]]] += staffed
    return days


def _write_month(ws, month: str, trial: list[tuple[int, int]] | None = None) -> None:
    ws.title = (("Probe " if trial is not None else "") + month_label(month))[:31]
    for h, c in COL.items():
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment = FONT_BOLD, FILL_HEADER, ALIGN_HEADER
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=MEDIUM)
    ws.row_dimensions[1].height = HEADER_HEIGHT

    data = _day_data(month, trial)
    _, m = parse_month(month)
    row = 2
    for d in (d for w in month_weeks(month) for d in w if d.month == m):
        iso = d.isoformat()
        day = data.get(iso, {"events": {}, "tageskasse": []})
        start = row
        ptr = start
        for ev in sorted(day["events"].values(), key=lambda e: (e["time"] or "99", e["title"])):
            ws.cell(row=ptr, column=COL["Uhrzeit"], value=ev["time"] or None)
            ws.cell(row=ptr, column=COL["Vorstellungen/Bühne"], value=ev["title"])
            height = max(1, *(len(ev[c]) for c in EVENT_COLUMNS))
            for col in EVENT_COLUMNS:
                for i, name in enumerate(ev[col]):
                    ws.cell(row=ptr + i, column=COL[col], value=name)
            ptr += height
        for i, name in enumerate(day["tageskasse"]):
            ws.cell(row=start + i, column=COL["Tageskasse"], value=name)
        block = max(MIN_ROWS_PER_DAY, ptr - start, len(day["tageskasse"]))

        date_cell = ws.cell(row=start, column=COL["Datum"], value=date.fromisoformat(iso))
        date_cell.number_format = "DD.MM.YYYY"
        ws.cell(row=start, column=COL["Tag"], value=WEEKDAYS[d.weekday()])

        for r in range(start, start + block):
            ws.row_dimensions[r].height = ROW_HEIGHT
            for h, c in COL.items():
                cell = ws.cell(row=r, column=c)
                cell.font = FONT_OPEN if cell.value == OPEN else (
                    FONT_BOLD if h == "Vorstellungen/Bühne" and cell.value else FONT)
                cell.fill = FILL_APP if h in APP_COLUMNS else FILL_FREE
                cell.alignment = ALIGN
                cell.border = Border(left=THIN, right=THIN, top=MEDIUM if r == start else THIN, bottom=THIN)
        row = start + block

    # Abschlusslinie + Legende
    for c in COL.values():
        ws.cell(row=row - 1, column=c).border = Border(
            left=THIN, right=THIN, top=ws.cell(row=row - 1, column=c).border.top, bottom=MEDIUM)
    legend = ws.cell(row=row + 1, column=1,
                     value=("PROBEPLAN – nicht veröffentlicht, Angaben können sich noch ändern · "
                            if trial is not None else "") + "Hellblau = aus dem Schichtplaner (Änderungen bitte in der App vornehmen und neu "
                           "exportieren) · Weiß = frei ausfüllbar (ASL, Proben, Sonstiges) · "
                           "„offen“ = noch unbesetzt")
    legend.font = FONT_LEGEND
    legend.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
    ws.row_dimensions[row].height = ROW_HEIGHT
    ws.row_dimensions[row + 1].height = ROW_HEIGHT

    for h, c in COL.items():
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[h]   # explizit setzen
    ws.freeze_panes = "C2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def export_regieplan(months: list[str], trial: dict[str, list[tuple[int, int]]] | None = None) -> bytes:
    """Excel-Datei mit einem Tabellenblatt je Monat. trial = {Monat: Probe-Zuteilung} für einen Probeplan."""
    wb = Workbook()
    wb.remove(wb.active)
    for month in months:
        _write_month(wb.create_sheet(), month, (trial or {}).get(month))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def months_with_shifts(around: str | None = None) -> list[str]:
    """Alle Monate, für die es Schichten gibt (eine einzige Abfrage)."""
    return db.months_with_shifts()
